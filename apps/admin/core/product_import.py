# ============================================================
# TNL-PRODUCT-IMPORT-XLSX-V1
# ============================================================

from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import os
import tempfile
import time
import uuid
import zipfile

from django import forms
from django.core import signing
from django.db import IntegrityError, transaction

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Catalogo, PerfilUsuario, Producto


HEADERS = [
    "sku",
    "nombre",
    "descripcion",
    "precio",
    "stock",
    "unidad",
    "identificador_externo",
    "activo",
]

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_PRODUCTS = 5000
MAX_ERRORS = 200

MAX_ZIP_ENTRIES = 2000
MAX_ZIP_TOTAL = 60 * 1024 * 1024
MAX_ZIP_MEMBER = 25 * 1024 * 1024

SIGNING_SALT = "tnl-product-import-xlsx-v1"
SIGNING_MAX_AGE = 60 * 60

PRIVATE_DIR = (
    Path(tempfile.gettempdir())
    / "tnl-product-imports"
)


class ProductoImportError(Exception):
    pass


class ProductoImportarExcelForm(forms.Form):

    catalogo = forms.ModelChoiceField(
        queryset=Catalogo.objects.none(),
        label="Catálogo destino",
        empty_label="Selecciona un catálogo",
        widget=forms.Select(
            attrs={
                "class": "form-control",
            }
        ),
    )

    archivo = forms.FileField(
        label="Archivo Excel (.xlsx)",
        widget=forms.ClearableFileInput(
            attrs={
                "class": "form-control",
                "accept": ".xlsx",
            }
        ),
    )

    def __init__(
        self,
        *args,
        catalogos=None,
        **kwargs,
    ):

        super().__init__(*args, **kwargs)

        self.fields[
            "catalogo"
        ].queryset = (
            catalogos
            if catalogos is not None
            else Catalogo.objects.none()
        )

    def clean_archivo(self):

        archivo = self.cleaned_data["archivo"]

        nombre = str(
            getattr(
                archivo,
                "name",
                "",
            )
        ).strip().lower()

        if not nombre.endswith(".xlsx"):
            raise forms.ValidationError(
                "Solo se permiten archivos .xlsx."
            )

        if archivo.size <= 0:
            raise forms.ValidationError(
                "El archivo está vacío."
            )

        if archivo.size > MAX_FILE_SIZE:
            raise forms.ValidationError(
                "El archivo Excel no puede "
                "superar 5 MB."
            )

        return archivo


def catalogos_autorizados_producto(user):

    queryset = (
        Catalogo.objects
        .select_related("empresa")
        .filter(activo=True)
        .order_by(
            "empresa__nombre",
            "nombre",
        )
    )

    perfil = getattr(
        user,
        "perfil_negociolisto",
        None,
    )

    if perfil and perfil.activo:

        if (
            perfil.rol
            == PerfilUsuario.Rol.CLIENTE
        ):

            if perfil.empresa_id is None:
                return queryset.none()

            return queryset.filter(
                empresa_id=perfil.empresa_id
            )

        return queryset

    if user.is_superuser or user.is_staff:
        return queryset

    return queryset.none()


def _read_bytes(source):

    if isinstance(
        source,
        (str, os.PathLike, Path),
    ):

        path = Path(source)

        if not path.is_file():
            raise ProductoImportError(
                "El archivo temporal "
                "ya no existe."
            )

        size = path.stat().st_size

        if (
            size <= 0
            or size > MAX_FILE_SIZE
        ):
            raise ProductoImportError(
                "Tamaño de archivo inválido."
            )

        return path.read_bytes()

    try:
        source.seek(0)
    except Exception:
        pass

    data = source.read()

    if (
        not data
        or len(data) > MAX_FILE_SIZE
    ):
        raise ProductoImportError(
            "Tamaño de archivo inválido."
        )

    return data


def _validar_zip(data):

    try:

        with zipfile.ZipFile(
            BytesIO(data)
        ) as archive:

            members = archive.infolist()

            if len(members) > MAX_ZIP_ENTRIES:
                raise ProductoImportError(
                    "El Excel contiene demasiados "
                    "elementos internos."
                )

            total = 0

            for member in members:

                if member.file_size > MAX_ZIP_MEMBER:
                    raise ProductoImportError(
                        "El Excel contiene un "
                        "elemento interno "
                        "demasiado grande."
                    )

                total += member.file_size

                if total > MAX_ZIP_TOTAL:
                    raise ProductoImportError(
                        "El contenido descomprimido "
                        "del Excel es demasiado "
                        "grande."
                    )

    except zipfile.BadZipFile as exc:

        raise ProductoImportError(
            "El archivo no es un .xlsx válido."
        ) from exc


def _texto(
    value,
    *,
    field,
    max_length=None,
    required=False,
):

    if value is None:
        text = ""

    elif isinstance(value, bool):
        text = str(value)

    elif isinstance(value, int):
        text = str(value)

    elif isinstance(value, float):

        if value.is_integer():
            text = str(int(value))
        else:
            text = str(value)

    else:
        text = str(value)

    text = text.strip()

    if required and not text:
        raise ValueError(
            f"{field}: valor obligatorio."
        )

    if (
        max_length is not None
        and len(text) > max_length
    ):
        raise ValueError(
            f"{field}: máximo "
            f"{max_length} caracteres."
        )

    return text


def _decimal(value, *, field):

    if value is None or value == "":
        raise ValueError(
            f"{field}: valor obligatorio."
        )

    if isinstance(value, bool):
        raise ValueError(
            f"{field}: debe ser numérico."
        )

    if isinstance(value, Decimal):
        number = value

    elif isinstance(
        value,
        (int, float),
    ):
        number = Decimal(str(value))

    else:

        text = str(value).strip()

        if (
            "," in text
            and "." in text
        ):
            raise ValueError(
                f"{field}: usa un solo "
                "separador decimal."
            )

        if (
            "," in text
            and "." not in text
        ):
            text = text.replace(",", ".")

        try:
            number = Decimal(text)

        except InvalidOperation as exc:
            raise ValueError(
                f"{field}: valor numérico "
                "inválido."
            ) from exc

    if not number.is_finite():
        raise ValueError(
            f"{field}: valor numérico inválido."
        )

    if number < 0:
        raise ValueError(
            f"{field}: no puede ser negativo."
        )

    normalized = (
        number.normalize()
        if number != 0
        else Decimal("0")
    )

    decimals = max(
        -normalized.as_tuple().exponent,
        0,
    )

    if decimals > 4:
        raise ValueError(
            f"{field}: máximo 4 decimales."
        )

    integer_digits = (
        max(
            normalized.adjusted() + 1,
            1,
        )
        if normalized != 0
        else 1
    )

    if integer_digits > 10:
        raise ValueError(
            f"{field}: máximo 10 "
            "dígitos enteros."
        )

    return number.quantize(
        Decimal("0.0001")
    )


def _booleano(value):

    if value is None:
        return True

    if isinstance(value, bool):
        return value

    if (
        isinstance(value, int)
        and value in (0, 1)
    ):
        return bool(value)

    text = str(value).strip().casefold()

    if not text:
        return True

    if text in {
        "1",
        "si",
        "sí",
        "s",
        "true",
        "verdadero",
        "yes",
        "y",
        "activo",
    }:
        return True

    if text in {
        "0",
        "no",
        "n",
        "false",
        "falso",
        "inactivo",
    }:
        return False

    raise ValueError(
        "activo: usa SI/NO, "
        "TRUE/FALSE o 1/0."
    )


def _add_error(errors, message):

    if len(errors) < MAX_ERRORS:
        errors.append(message)


def validar_archivo_productos(
    source,
    *,
    catalogo,
):

    global_errors = []
    row_errors = []
    rows = []

    try:

        data = _read_bytes(source)

        _validar_zip(data)

        workbook = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=False,
        )

    except ProductoImportError as exc:

        return {
            "ok": False,
            "errores_globales": [str(exc)],
            "errores_filas": [],
            "filas": [],
        }

    except Exception:

        return {
            "ok": False,
            "errores_globales": [
                "No se pudo abrir el archivo. "
                "Confirma que sea un .xlsx válido."
            ],
            "errores_filas": [],
            "filas": [],
        }

    try:

        if "Productos" not in workbook.sheetnames:

            return {
                "ok": False,
                "errores_globales": [
                    "El archivo debe contener una "
                    'hoja llamada "Productos".'
                ],
                "errores_filas": [],
                "filas": [],
            }

        worksheet = workbook["Productos"]

        if worksheet.max_column > 20:
            global_errors.append(
                "El archivo contiene "
                "demasiadas columnas."
            )

        if (
            worksheet.max_row
            > MAX_PRODUCTS + 1
        ):
            global_errors.append(
                "El archivo supera el máximo "
                f"de {MAX_PRODUCTS} productos."
            )

        if global_errors:

            return {
                "ok": False,
                "errores_globales": global_errors,
                "errores_filas": [],
                "filas": [],
            }

        first_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=False,
            ),
            (),
        )

        raw_headers = [
            (
                str(cell.value)
                .strip()
                .casefold()
                if cell.value is not None
                else ""
            )
            for cell in first_row
        ]

        while (
            raw_headers
            and raw_headers[-1] == ""
        ):
            raw_headers.pop()

        headers = [
            x for x in raw_headers if x
        ]

        duplicates = sorted(
            {
                x
                for x in headers
                if headers.count(x) > 1
            }
        )

        missing = [
            x
            for x in HEADERS
            if x not in headers
        ]

        extra = [
            x
            for x in headers
            if x not in HEADERS
        ]

        if duplicates:
            global_errors.append(
                "Encabezados duplicados: "
                + ", ".join(duplicates)
                + "."
            )

        if missing:
            global_errors.append(
                "Faltan columnas: "
                + ", ".join(missing)
                + "."
            )

        if extra:
            global_errors.append(
                "Columnas no reconocidas: "
                + ", ".join(extra)
                + "."
            )

        if global_errors:

            return {
                "ok": False,
                "errores_globales": global_errors,
                "errores_filas": [],
                "filas": [],
            }

        index = {
            header: raw_headers.index(header)
            for header in HEADERS
        }

        for excel_row, cells in enumerate(
            worksheet.iter_rows(
                min_row=2,
                max_col=len(raw_headers),
                values_only=False,
            ),
            start=2,
        ):

            values = [
                cell.value
                for cell in cells
            ]

            if all(
                value is None
                or (
                    isinstance(value, str)
                    and not value.strip()
                )
                for value in values
            ):
                continue

            formula_fields = []

            for header in HEADERS:

                pos = index[header]

                if (
                    pos < len(cells)
                    and cells[pos].data_type
                    == "f"
                ):
                    formula_fields.append(
                        header
                    )

            if formula_fields:

                _add_error(
                    row_errors,
                    f"Fila {excel_row}: "
                    "no se permiten fórmulas "
                    "en: "
                    + ", ".join(
                        formula_fields
                    )
                    + "."
                )

                continue

            try:

                # TNL-PRODUCT-SKU-OPTIONAL-REPEATABLE-V1
                sku = _texto(
                    values[index["sku"]],
                    field="sku",
                    max_length=100,
                    required=False,
                )

                nombre = _texto(
                    values[index["nombre"]],
                    field="nombre",
                    max_length=200,
                    required=True,
                )

                descripcion = _texto(
                    values[
                        index["descripcion"]
                    ],
                    field="descripcion",
                )

                precio = _decimal(
                    values[index["precio"]],
                    field="precio",
                )

                stock = _decimal(
                    values[index["stock"]],
                    field="stock",
                )

                unidad = _texto(
                    values[index["unidad"]],
                    field="unidad",
                    max_length=50,
                    required=True,
                )

                externo = _texto(
                    values[
                        index[
                            "identificador_externo"
                        ]
                    ],
                    field=(
                        "identificador_externo"
                    ),
                    max_length=255,
                )

                activo = _booleano(
                    values[index["activo"]]
                )

            except ValueError as exc:

                _add_error(
                    row_errors,
                    f"Fila {excel_row}: {exc}"
                )

                continue

            # SKU es informativo y puede repetirse.
            # Cada fila representa un producto independiente.

            rows.append(
                {
                    "fila": excel_row,
                    "sku": sku,
                    "nombre": nombre,
                    "descripcion": descripcion,
                    "precio": str(precio),
                    "stock": str(stock),
                    "unidad": unidad,
                    "identificador_externo": (
                        externo
                    ),
                    "activo": activo,
                }
            )

        if not rows and not row_errors:

            global_errors.append(
                "El archivo no contiene "
                "productos."
            )

        # No se valida unicidad del SKU.
        # Puede estar vacío o repetirse en el catálogo.

        if len(row_errors) >= MAX_ERRORS:

            global_errors.append(
                "Se alcanzó el límite de "
                f"{MAX_ERRORS} errores."
            )

        ok = (
            not global_errors
            and not row_errors
            and bool(rows)
        )

        return {
            "ok": ok,
            "errores_globales": global_errors,
            "errores_filas": row_errors,
            "filas": (
                rows
                if ok
                else []
            ),
        }

    finally:
        workbook.close()


def _private_dir():

    PRIVATE_DIR.mkdir(
        mode=0o700,
        parents=True,
        exist_ok=True,
    )

    try:
        PRIVATE_DIR.chmod(0o700)
    except OSError:
        pass

    return PRIVATE_DIR


def limpiar_temporales():

    directory = _private_dir()

    cutoff = (
        time.time()
        - (SIGNING_MAX_AGE * 2)
    )

    for path in directory.glob("*.xlsx"):

        try:

            if (
                path.is_file()
                and path.stat().st_mtime
                < cutoff
            ):
                path.unlink()

        except OSError:
            pass


def guardar_importacion_temporal(archivo):

    limpiar_temporales()

    directory = _private_dir()

    import_id = uuid.uuid4().hex

    path = (
        directory
        / f"{import_id}.xlsx"
    )

    digest = sha256()
    total = 0

    try:
        archivo.seek(0)
    except Exception:
        pass

    try:

        with open(path, "xb") as destination:

            os.chmod(path, 0o600)

            for chunk in archivo.chunks():

                total += len(chunk)

                if total > MAX_FILE_SIZE:
                    raise ProductoImportError(
                        "El archivo supera 5 MB."
                    )

                digest.update(chunk)
                destination.write(chunk)

        if total <= 0:
            raise ProductoImportError(
                "El archivo está vacío."
            )

    except Exception:

        try:
            path.unlink(missing_ok=True)
        except OSError:
            pass

        raise

    return {
        "import_id": import_id,
        "sha256": digest.hexdigest(),
        "size": total,
        "path": path,
    }


def eliminar_importacion_temporal(
    import_id,
):

    try:

        normalized = uuid.UUID(
            str(import_id)
        ).hex

    except (
        ValueError,
        TypeError,
        AttributeError,
    ):
        return False

    path = (
        _private_dir()
        / f"{normalized}.xlsx"
    )

    if path.is_file():

        path.unlink()

        return True

    return False


def firmar_importacion(
    *,
    catalogo_id,
    import_id,
    file_hash,
    row_count,
):

    return signing.dumps(
        {
            "catalogo_id": int(
                catalogo_id
            ),
            "import_id": import_id,
            "sha256": file_hash,
            "row_count": int(
                row_count
            ),
        },
        salt=SIGNING_SALT,
        compress=True,
    )


def leer_importacion_firmada(token):

    return signing.loads(
        token,
        salt=SIGNING_SALT,
        max_age=SIGNING_MAX_AGE,
    )


def resolver_importacion_temporal(
    payload,
):

    try:

        import_id = uuid.UUID(
            str(payload["import_id"])
        ).hex

        expected_hash = str(
            payload["sha256"]
        )

        expected_rows = int(
            payload["row_count"]
        )

    except (
        KeyError,
        ValueError,
        TypeError,
        AttributeError,
    ) as exc:

        raise ProductoImportError(
            "Token de importación inválido."
        ) from exc

    if (
        len(expected_hash) != 64
        or expected_rows < 1
        or expected_rows > MAX_PRODUCTS
    ):
        raise ProductoImportError(
            "Token de importación inválido."
        )

    path = (
        _private_dir()
        / f"{import_id}.xlsx"
    )

    if not path.is_file():
        raise ProductoImportError(
            "El archivo temporal expiró "
            "o ya no existe."
        )

    size = path.stat().st_size

    if (
        size <= 0
        or size > MAX_FILE_SIZE
    ):
        raise ProductoImportError(
            "Archivo temporal inválido."
        )

    digest = sha256(
        path.read_bytes()
    ).hexdigest()

    if digest != expected_hash:
        raise ProductoImportError(
            "El archivo temporal cambió "
            "después de validarlo."
        )

    return (
        import_id,
        path,
        expected_rows,
    )


def importar_productos_validados(
    *,
    catalogo,
    rows,
):

    if (
        not isinstance(rows, list)
        or not rows
        or len(rows) > MAX_PRODUCTS
    ):
        raise ProductoImportError(
            "Lote de importación inválido."
        )

    # SKU no identifica al registro.
    # Cada fila validada se crea como producto independiente.

    try:

        with transaction.atomic():

            locked_catalog = (
                Catalogo.objects
                .select_for_update()
                .get(pk=catalogo.pk)
            )
            # TNL-RESTAURANTE-IMPORT-CATEGORIA-V1
            categoria_restaurante_default = None

            plantilla_catalogo = getattr(
                locked_catalog,
                "plantilla",
                None,
            )

            if (
                str(
                    getattr(
                        plantilla_catalogo,
                        "tipo",
                        "",
                    )
                    or ""
                )
                == "restaurante"
            ):

                from .models import CategoriaProducto

                categoria_restaurante_default = (
                    CategoriaProducto.objects
                    .filter(
                        catalogo=locked_catalog,
                        nombre="Menú general",
                        activa=True,
                    )
                    .first()
                )

                if categoria_restaurante_default is None:

                    raise ProductoImportError(
                        "El catálogo Restaurante no tiene "
                        "una categoría general activa."
                    )


            # No se consulta existencia por SKU.
            # Se permiten valores duplicados y productos sin SKU.

            products = []

            for row in rows:

                products.append(
                    Producto(
                        catalogo=locked_catalog,
                        categoria=categoria_restaurante_default,
                        sku=row["sku"],
                        nombre=row["nombre"],
                        descripcion=row.get(
                            "descripcion",
                            "",
                        ),
                        precio=Decimal(
                            row["precio"]
                        ),
                        stock=Decimal(
                            row["stock"]
                        ),
                        unidad=row["unidad"],
                        identificador_externo=(
                            row.get(
                                "identificador_externo",
                                "",
                            )
                        ),
                        activo=bool(
                            row.get(
                                "activo",
                                True,
                            )
                        ),
                    )
                )

            Producto.objects.bulk_create(
                products,
                batch_size=500,
            )

            return len(products)

    except IntegrityError as exc:

        raise ProductoImportError(
            "La base de datos rechazó "
            "el lote. No se creó ningún "
            "producto."
        ) from exc


def generar_plantilla_productos_xlsx():

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Productos"

    fill = PatternFill(
        "solid",
        fgColor="1F2937",
    )

    font = Font(
        color="FFFFFF",
        bold=True,
    )

    for column, header in enumerate(
        HEADERS,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column,
            value=header,
        )

        cell.fill = fill
        cell.font = font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    widths = {
        "A": 22,
        "B": 34,
        "C": 48,
        "D": 16,
        "E": 16,
        "F": 18,
        "G": 28,
        "H": 14,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:H1"

    validation = DataValidation(
        type="list",
        formula1='"SI,NO"',
        allow_blank=True,
    )

    worksheet.add_data_validation(
        validation
    )

    validation.add("H2:H5001")

    instructions = workbook.create_sheet(
        "Instrucciones"
    )

    instructions.append(
        [
            "Campo",
            "Obligatorio",
            "Descripción",
        ]
    )

    rows = [
        (
            "sku",
            "No",
            "Opcional. Puede contener texto, números "
            "o ambos y puede repetirse.",
        ),
        (
            "nombre",
            "Sí",
            "Nombre del producto.",
        ),
        (
            "descripcion",
            "No",
            "Descripción comercial.",
        ),
        (
            "precio",
            "Sí",
            "Mayor o igual a cero. "
            "Máximo 4 decimales.",
        ),
        (
            "stock",
            "Sí",
            "Mayor o igual a cero. "
            "Máximo 4 decimales.",
        ),
        (
            "unidad",
            "Sí",
            "Ej. pieza, caja, kg, litro.",
        ),
        (
            "identificador_externo",
            "No",
            "ID en otro sistema.",
        ),
        (
            "activo",
            "No",
            "SI o NO. Vacío = SI.",
        ),
    ]

    for row in rows:
        instructions.append(row)

    instructions.append([])

    instructions.append(
        [
            "IMPORTANTE",
            "",
            "Empresa y catálogo no se "
            "capturan en el Excel. "
            "El destino se selecciona "
            "en NegocioListo.",
        ]
    )

    instructions.append(
        [
            "IMÁGENES",
            "",
            "Las imágenes se cargan desde "
            "la ficha individual del "
            "producto. Máximo 4 MB "
            "por imagen.",
        ]
    )

    for cell in instructions[1]:
        cell.fill = fill
        cell.font = font

    instructions.column_dimensions[
        "A"
    ].width = 28

    instructions.column_dimensions[
        "B"
    ].width = 16

    instructions.column_dimensions[
        "C"
    ].width = 82

    stream = BytesIO()

    workbook.save(stream)
    workbook.close()

    return stream.getvalue()
